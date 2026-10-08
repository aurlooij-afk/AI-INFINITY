UPLOAD_MAX_BYTES = max(5_000_000, int(os.getenv("AI_INFINITY_MAX_UPLOAD_BYTES","100_000_000")))

def _safe_upload_name(name: str) -> str:
    base=safe_name(Path(str(name or "upload")).name)
    return (base or "upload")[:140]

def _extract_source_text(path: Path) -> str:
    """Extract bounded source text without making a successful upload depend on one parser."""
    ext=path.suffix.lower()
    try:
        if ext in {".txt",".md",".json",".csv",".srt",".vtt",".html",".htm"}:
            raw=path.read_text(encoding="utf-8",errors="ignore")
            if ext in {".html",".htm"}:
                raw=re.sub(r"(?is)<script[^>]*>.*?</script>|<style[^>]*>.*?</style>"," ",raw)
                raw=re.sub(r"(?s)<[^>]+>"," ",raw)
                raw=html.unescape(raw)
            return re.sub(r"\s+"," ",raw).strip()[:200_000]
        if ext==".pdf":
            try:
                from pypdf import PdfReader
                reader=PdfReader(str(path))
                text="\n".join((page.extract_text() or "") for page in reader.pages[:100])
                if text.strip():
                    return text[:200_000]
            except Exception:
                pass
            if shutil.which("pdftotext"):
                proc=subprocess.run(["pdftotext","-layout",str(path),"-"],capture_output=True,text=True,timeout=45)
                if proc.returncode==0 and proc.stdout.strip():
                    return proc.stdout[:200_000]
        if ext==".docx":
            from docx import Document
            doc=Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs if p.text)[:200_000]
        if ext==".pptx":
            try:
                from pptx import Presentation
                prs=Presentation(str(path))
                slides=[]
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if getattr(shape,"has_text_frame",False):
                            slides.append(shape.text)
                return "\n".join(x for x in slides if x)[:200_000]
            except Exception:
                return ""
    except Exception:
        return ""
    return ""

def _gap_audit_3624(user_id: str) -> Dict[str, Any]:
    # 60 domains x 20 acceptance controls = 1,200 concrete checks. These are
    # acceptance controls, not decorative "feature count" inflation.
    domains=[
        "command","briefing","research","evidence","claims","story","hooks","script","editorial","visual-direction",
        "image","video","voice","music","sfx","captions","accessibility","localization","rtl","thumbnail",
        "shorts","repurpose","timeline","preview","assets","library","brand","templates","projects","queue",
        "progress","cancel","retry","resume","quality-control","provenance","licensing","package","download","publishing",
        "scheduling","analytics","trends","audience","seo","social-copy","podcast","workspace","connections","agents",
        "security","privacy","performance","mobile","persistence","backup","observability","truthfulness","provider-routing","extensibility"
    ]
    controls=[
        "request validation","input limits","clear status","real artifact evidence","persistent metadata",